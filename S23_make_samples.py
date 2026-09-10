import docx
from reportlab.pdfgen import canvas

# 1. 生成一个中文 Word 测试文件
doc=docx.Document()
doc.add_paragraph("虚拟环境是 Python 用来隔离项目依赖的工具。")
doc.add_paragraph("它让不同项目可以安装不同版本的包，互不干扰。")
doc.save("sample.docx")

# 2. 生成一个 PDF 测试文件（英文，避免中文字体问题）
c=canvas.Canvas("sample.pdf")
c.drawString(100,700,"Python virtual environment test.")
c.drawString(100,800,"This is a sample PDF for parsing")
c.save()

print("测试文件已生成: sample.docx,sample.pdf")
