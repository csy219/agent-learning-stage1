import re
from pathlib import Path

import docx
from pypdf import PdfReader

#- PDF 每一页提取出来的文字经常有多余换行、多个空格、页眉页码；
# - \s+ 表示“一个或多个空白（空格/换行/Tab）”；
# - re.sub(r"\s+", " ", raw) = 把所有连续空白都替换成一个空格；
# - .strip() = 去掉开头结尾空白。
# 意义：PDF 文字往往很脏，RAG 里“脏数据 → 干净数据”就靠这种清洗函数。
def clean_text(raw:str)->str:
    """把多行文字合并成干净的连续文本"""
    return re.sub(r"\s+"," ",raw).strip()


# reader.pages = PDF 的所有页；
# page.extract_text() = 提取这一页的文字；
# or "" = 如果某页提取不出文字（扫描图），给空字符串，避免程序崩溃；
# "\n".join(pages) = 把每一页文字用换行连起来；
# 最后交给 clean_text 清洗。
def extract_pdf(path:Path)->str:
    """提取PDF每一页的文字"""
    reader=PdfReader(str(path))
    pages=[page.extract_text() or ""for page in reader.pages]
    return clean_text("\n".join(pages))



# Word 的本质是一堆“段落”，document.paragraphs 就是所有段落；
# p.text = 取这个段落的文字；
# 列表推导式 [p.text for p in document.paragraphs] = “把每个段落都取出来，装进列表”。
def extract_docx(path:Path)->str:
    """提取word文档里的段落文字"""
    document=docx.Document(str(path))
    paragraphs=[p.text for p in document.paragraphs]
    return clean_text("\n".join(paragraphs))

if __name__ =="__main__":
    pdf_path=Path("sample.pdf")
    docx_path=Path("sample.docx")

    if pdf_path.exists():
        print("PDF 内容: ",extract_pdf(pdf_path))

    if docx_path.exists():
        print("Word 内容: ",extract_docx(docx_path))