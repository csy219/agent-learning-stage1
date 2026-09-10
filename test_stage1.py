from pathlib import Path

import docx
import pytest
from reportlab.pdfgen import canvas
from pydantic import ValidationError


from parse_docs import clean_text,extract_docx,extract_pdf
from validation_utils import validate_contact

# ---------- clean_text 的测试 ----------
def test_clean_text_collapses_spaces():
    assert clean_text("hello   world")=="hello world"

def test_clean_text_collapses_newlines():
    assert clean_text("hello\n\nworld")=="hello world"

def test_clean_text_strips():
    assert clean_text("    helloworld    ")=="helloworld"

# ---------- Pydantic 校验的测试 ----------
def test_validate_contact_ok():
    contact=validate_contact(
            '{"name": "张三", "phone": "13812345678", "date": "2026-01-01"}'
    )
    assert contact.name=="张三"
    assert contact.phone=="13812345678"

def test_validate_contact_missing_phone():
    with pytest.raises(ValidationError):
        validate_contact('{"name":"张三","date":""2026-01-01}')

def test_validate_contact_missing_date():
    with pytest.raises(ValidationError):
        validate_contact('{"name": "张三", "phone": "13812345678"}')

# ---------- 文档解析的测试 ----------
def test_extract_docx(tmp_path:Path):
    path=tmp_path/"sample.docx"
    document=docx.Document()
    document.add_paragraph("第一段内容")
    document.add_paragraph("第二段内容")
    document.save(str(path))

    text=extract_docx(path)
    assert "第一段内容" in text
    assert "第二段内容" in text

def test_extract_pdf(tmp_path:Path):
    path=tmp_path/"sample.pdf"
    pdf=canvas.Canvas(str(path))
    pdf.drawString(100,700,"Hello PDF")
    pdf.save()

    text=extract_pdf(path)
    assert "Hello PDF" in text